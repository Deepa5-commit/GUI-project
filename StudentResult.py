import tkinter as tk
from tkinter import messagebox
import openpyxl
from openpyxl import Workbook, load_workbook
import os

window=tk.Tk()
window.title("Student Result Management System")
window.geometry("800x500")
window.config(bg="lightblue")

file="student_results.xlsx"

if not os.path.exists(file):
    wb= Workbook()
    ws= wb.active
    ws.title="Student Result"
    ws.append(["Name","Roll no.","Class","Sub1 marks","Sub2 marks","Sub3 marks","Sub4 marks","Sub5 marks","Total marks","Percentage","Result"])
    wb.save(file)

def add_student():
    frame=tk.Frame(window)
    frame.pack()
    
    l1=tk.Label(frame,text="Name:")
    l1.grid(row=1,column=0,padx=10,pady=10)
    name_entry=tk.Entry(frame)
    name_entry.grid(row=1,column=1)
    
    l2=tk.Label(frame,text="Roll no. :")
    l2.grid(row=2,column=0,padx=10,pady=10)
    rollno_entry=tk.Entry(frame)
    rollno_entry.grid(row=2,column=1)

    l3=tk.Label(frame,text="Class:")
    l3.grid(row=3,column=0,padx=10,pady=10)
    stdclass_entry=tk.Entry(frame)
    stdclass_entry.grid(row=3,column=1)
    
    l4=tk.Label(frame,text="Subject 1 marks:")
    l4.grid(row=4,column=0,padx=10,pady=10)
    sub1_entry=tk.Entry(frame)
    sub1_entry.grid(row=4,column=1)
    
    l5=tk.Label(frame,text="Subject 2 marks:")
    l5.grid(row=5,column=0,padx=10,pady=10)
    sub2_entry=tk.Entry(frame)
    sub2_entry.grid(row=5,column=1)
    
    l6=tk.Label(frame,text="Subject 3 marks:")
    l6.grid(row=6,column=0,padx=10,pady=10)
    sub3_entry=tk.Entry(frame)
    sub3_entry.grid(row=6,column=1)
    
    l7=tk.Label(frame,text="Subject 4 marks:")
    l7.grid(row=7,column=0,padx=10,pady=10)
    sub4_entry=tk.Entry(frame)
    sub4_entry.grid(row=7,column=1)
    
    l8=tk.Label(frame,text="Subject 5 marks:")
    l8.grid(row=8,column=0,padx=10,pady=10)
    sub5_entry=tk.Entry(frame)
    sub5_entry.grid(row=8,column=1)

    def save_student():
        name=name_entry.get()
        rollno=rollno_entry.get()
        stdclass=stdclass_entry.get()
        try:
            sub1=int(sub1_entry.get())
            sub2=int(sub2_entry.get())
            sub3=int(sub3_entry.get())
            sub4=int(sub4_entry.get())
            sub5=int(sub5_entry.get())
        except ValueError:
            messagebox.showerror("Error,Enter valid marks")
            return
        tot_marks=sub1+sub2+sub3+sub4+sub5
        perc=(tot_marks/500)*100
        if sub1>=35 and sub2>=35 and sub3>=35 and sub4>=35 and sub5>=35:
            result="PASS"
        else:
            result="FAIL"

        wb=load_workbook(file)
        ws=wb.active
        ws.append([name,rollno,stdclass,sub1,sub2,sub3,sub4,sub5,tot_marks,perc,result])
        wb.save(file)
        messagebox.showinfo("Success, Data saved in Excel")

    submit=tk.Button(frame,text="Save",command=save_student)
    submit.grid(row=10,column=0,columnspan=2, pady=15)     

def get_result():
    frame=tk.Frame(window)
    frame.pack()
    
    l9=tk.Label(frame, text="Enter Roll no.")
    l9.grid(row=12,column=0,padx=10,pady=10)
    rollno=tk.Entry(frame)
    rollno.grid(row=12,column=1)

    def search():
        wb=load_workbook(file)
        ws=wb.active
        for row in ws.iter_rows(min_row=2,values_only=True):
            if str(row[1])==rollno.get():
                messagebox.showinfo("Student Result")
                                
                return
        messagebox.showerror("Error , Roll no. not found")

    bt2=tk.Button(frame, text="Search", command=search)
    bt2.grid(row=13,column=1,pady=10)
        
def show_results():
    wb=load_workbook(file)
    ws=wb.active
    text=""
    for row in ws.iter_rows(min_row=2,values_only=True):
        text=text+ str(row)+ "\n\n"
    messagebox.showinfo("All Results",text)

def exit():
    window.destroy()
    
label=tk.Label(window, text="Student Result Management", font=("Arial",30,"bold")).pack(pady=25)
    
btn1=tk.Button(window,text="Add Student",command=add_student)
btn1.pack(pady=15)

btn2=tk.Button(window,text="Get Result",command=get_result)
btn2.pack(pady=15)

btn3=tk.Button(window,text="Show All Results",command=show_results)
btn3.pack(pady=15)

btn4=tk.Button(window,text="Exit",command=exit)
btn4.pack(pady=15)

window.mainloop()
